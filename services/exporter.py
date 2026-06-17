"""Export a TPA invoice to XLSX."""
import io
from datetime import date

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

INDIGO = "01325C"   # Oakmore Labs navy
PRODUCT = "Oakmore Labs · TPA Billing"


def build_xlsx(invoice, client, contract):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"

    bold = Font(bold=True)
    white_bold = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=INDIGO)
    money = '#,##0.00'
    border = Border(bottom=Side(style="thin", color="DDDDDD"))

    ws["A1"] = "INVOICE"
    ws["A1"].font = Font(bold=True, size=16, color=INDIGO)
    ws["A2"] = PRODUCT
    ws["A2"].font = Font(italic=True, color="666666")

    ws["A4"] = "Bill to:"
    ws["A4"].font = bold
    ws["A5"] = client.name
    ws["D4"] = "Invoice #:"
    ws["E4"] = invoice.number
    ws["D5"] = "Period:"
    ws["E5"] = invoice.month
    ws["D6"] = "Contract:"
    ws["E6"] = contract.number if contract else ""
    ws["D7"] = "Terms:"
    ws["E7"] = invoice.terms
    ws["D8"] = "Date:"
    ws["E8"] = date.today().isoformat()
    for c in ("D4", "D5", "D6", "D7", "D8"):
        ws[c].font = bold

    hdr = ["Service", "Basis", "Rate", "Qty", "Amount"]
    r = 10
    for c, label in enumerate(hdr, start=1):
        cell = ws.cell(row=r, column=c, value=label)
        cell.font = white_bold
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center" if c > 2 else "left")

    r += 1
    total = 0.0
    for ln in invoice.lines:
        if not ln.charged:
            continue
        ws.cell(row=r, column=1, value=ln.service)
        ws.cell(row=r, column=2, value=ln.basis_label)
        ws.cell(row=r, column=3, value=round(ln.contract_rate, 2)).number_format = money
        ws.cell(row=r, column=4, value=("—" if ln.basis == "FLAT" else ln.billed_count))
        ws.cell(row=r, column=5, value=round(ln.amount, 2)).number_format = money
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = border
        total += ln.amount
        r += 1

    ws.cell(row=r, column=4, value="Total due").font = bold
    tc = ws.cell(row=r, column=5, value=round(total, 2))
    tc.font = bold
    tc.number_format = money
    r += 2

    held = [l for l in invoice.lines if not l.charged]
    if held:
        ws.cell(row=r, column=1,
                value=f"Held off invoice ({len(held)}): " +
                      "; ".join(f"{l.service} — {l.recon_reason}" for l in held)
                ).font = Font(italic=True, color="DD2C51")
        r += 2

    ws.cell(row=r, column=1,
            value="Reconciled against the administrative services agreement and the "
                  "period's eligibility file.").font = Font(italic=True, size=9, color="888888")

    widths = [34, 16, 12, 10, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
